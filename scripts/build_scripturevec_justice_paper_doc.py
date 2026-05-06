from __future__ import annotations

import csv
import json
import math
import re
import shutil
import textwrap
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/seth/projects/scriptorium/projects/virtue-bench-2")
SOURCE_MD = Path("/Users/seth/Downloads/scripturevec_justice_paper.md")
PACKET = ROOT / "results/paper/scripturevec_justice"
KEY_DATA = PACKET / "key_data"
WRITING_PACKET = PACKET / "writing_packet"
OUT_DIR = PACKET / "paper_doc"
FIG_DIR = PACKET / "figures"
DOCX_OUT = OUT_DIR / "scripturevec_justice_paper_updated.docx"
MD_OUT = OUT_DIR / "scripturevec_justice_paper_updated.md"
REVIEW_OUT = OUT_DIR / "scripturevec_justice_paper_evaluation.md"
PAPER_MD_OUT = ROOT / "paper/search_out_a_matter_scripturevec_justice.md"
PAPER_DOCX_OUT = ROOT / "paper/search_out_a_matter_scripturevec_justice.docx"
PAPER_REVIEW_OUT = ROOT / "paper/scripturevec_justice_paper_evaluation.md"
EXPANDED_SUMMARY = (
    ROOT
    / "results/experiments/scripturevec14/"
    / "scripturevec14_qwen3_14b_confirmed_chapter_layerloc_expandedgrid2_v1_summary.json"
)

FONT_REG = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
FONT_ITALIC = "/System/Library/Fonts/Supplemental/Arial Italic.ttf"

NAVY = "#172033"
BLUE = "#2563eb"
LIGHT_BLUE = "#dbeafe"
TEAL = "#0f766e"
AMBER = "#b45309"
GREEN = "#15803d"
RED = "#b91c1c"
GRAY = "#64748b"
LIGHT_GRAY = "#e5e7eb"
PAPER = "#fbfbf8"
INK = "#1f2937"


def read_csv(name: str) -> list[dict[str, str]]:
    with (KEY_DATA / name).open(newline="") as f:
        return list(csv.DictReader(f))


def write_csv(name: str, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    with (KEY_DATA / name).open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def fmt_num(value: object, digits: int = 4) -> str:
    if value is None:
        return ""
    try:
        out = f"{float(value):.{digits}f}".rstrip("0").rstrip(".")
        return out if out else "0"
    except (TypeError, ValueError):
        return str(value)


def fmt_alpha(value: object) -> str:
    return str(int(float(value)))


def paired_text(paired: dict[str, object] | None) -> str:
    if not paired:
        return ""
    improve = int(paired.get("improve", 0))
    regress = int(paired.get("regress", 0))
    changes = int(paired.get("answer_changes", 0))
    return f"{changes} chg / {improve - regress:+d} net"


def target_metadata() -> dict[str, dict[str, str]]:
    rows = read_csv("chapter_confirmation_l40_survivors.csv")
    return {r["target"]: r for r in rows}


def refresh_localization_key_data() -> None:
    """Refresh paper-facing localization CSVs from the completed expanded grid."""
    if not EXPANDED_SUMMARY.exists():
        return
    data = json.loads(EXPANDED_SUMMARY.read_text())
    meta = target_metadata()

    cells = sorted(
        data["by_cell"],
        key=lambda r: (
            -int(r["paired_rescues"]),
            -float(r["mean_positive_delta"]),
            int(r["center"]),
            float(r["alpha"]),
        ),
    )
    cell_rows = [
        {
            "center_layer": int(r["center"]),
            "alpha": fmt_alpha(r["alpha"]),
            "rows": int(r["row_count"]),
            "accuracy_rescues": int(r["accuracy_rescues"]),
            "paired_rescues": int(r["paired_rescues"]),
            "mean_positive_delta": fmt_num(r["mean_positive_delta"]),
            "rescued_targets": ";".join(r["targets"]),
        }
        for r in cells
    ]
    write_csv(
        "layer_alpha_cells.csv",
        [
            "center_layer",
            "alpha",
            "rows",
            "accuracy_rescues",
            "paired_rescues",
            "mean_positive_delta",
            "rescued_targets",
        ],
        cell_rows,
    )

    expected = sorted(data["by_cell"], key=lambda r: (int(r["center"]), float(r["alpha"])))
    expected_rows = [
        {
            "center_layer": int(r["center"]),
            "alpha": fmt_alpha(r["alpha"]),
            "status": "completed",
            "paired_rescues": int(r["paired_rescues"]),
            "mean_positive_delta": fmt_num(r["mean_positive_delta"]),
            "rescued_targets": ";".join(r["targets"]),
        }
        for r in expected
    ]
    write_csv(
        "layer_alpha_expected_grid.csv",
        ["center_layer", "alpha", "status", "paired_rescues", "mean_positive_delta", "rescued_targets"],
        expected_rows,
    )

    target_rows = []
    for r in sorted(data["by_target"], key=lambda r: (-int(r["paired_rescues"]), r["target"])):
        m = meta[r["target"]]
        best = r["best_cell"]
        target_rows.append(
            {
                "target": r["target"],
                "book": m["book"],
                "chapter": m["chapter"],
                "reference": m["reference"],
                "paired_rescue_count": int(r["paired_rescues"]),
                "accuracy_rescue_count": int(r["accuracy_rescues"]),
                "best_center_layer": int(best["center"]),
                "best_alpha": fmt_alpha(best["alpha"]),
                "best_control_accuracy": fmt_num(best["control_accuracy"]),
                "best_positive_accuracy": fmt_num(best["positive_accuracy"]),
                "best_negative_accuracy": fmt_num(best["negative_accuracy"]),
                "best_null_accuracy": fmt_num(best["null_accuracy"]),
                "best_positive_paired": paired_text(best.get("positive_paired")),
            }
        )
    write_csv(
        "chapter_stability_by_localization.csv",
        [
            "target",
            "book",
            "chapter",
            "reference",
            "paired_rescue_count",
            "accuracy_rescue_count",
            "best_center_layer",
            "best_alpha",
            "best_control_accuracy",
            "best_positive_accuracy",
            "best_negative_accuracy",
            "best_null_accuracy",
            "best_positive_paired",
        ],
        target_rows,
    )

    detail_rows = []
    for r in sorted(data["rows"], key=lambda r: (int(r["center"]), float(r["alpha"]), r["target"])):
        m = meta[r["target"]]
        detail_rows.append(
            {
                "target": r["target"],
                "book": m["book"],
                "chapter": m["chapter"],
                "reference": m["reference"],
                "center_layer": int(r["center"]),
                "alpha": fmt_alpha(r["alpha"]),
                "control_accuracy": fmt_num(r["control_accuracy"]),
                "positive_accuracy": fmt_num(r["positive_accuracy"]),
                "positive_delta": fmt_num(r["positive_delta"]),
                "negative_accuracy": fmt_num(r["negative_accuracy"]),
                "null_accuracy": fmt_num(r["null_accuracy"]),
                "positive_paired": paired_text(r.get("positive_paired")),
                "paired_rescue": "yes" if r.get("is_paired_rescue") else "no",
            }
        )
    write_csv(
        "layer_alpha_target_rows.csv",
        [
            "target",
            "book",
            "chapter",
            "reference",
            "center_layer",
            "alpha",
            "control_accuracy",
            "positive_accuracy",
            "positive_delta",
            "negative_accuracy",
            "null_accuracy",
            "positive_paired",
            "paired_rescue",
        ],
        detail_rows,
    )

    cell_headers = [f"L{int(r['center'])}_a{int(float(r['alpha'])):03d}" for r in expected]
    rescue_lookup = {
        (r["target"], f"L{int(r['center'])}_a{int(float(r['alpha'])):03d}"): "1"
        for r in data["rows"]
        if r.get("is_paired_rescue")
    }
    matrix_rows = []
    for r in target_rows:
        row = {
            "target": r["target"],
            "book": r["book"],
            "chapter": r["chapter"],
            "reference": r["reference"],
        }
        for head in cell_headers:
            row[head] = rescue_lookup.get((r["target"], head), "0")
        matrix_rows.append(row)
    write_csv(
        "chapter_x_layer_alpha_rescue_matrix.csv",
        ["target", "book", "chapter", "reference"] + cell_headers,
        matrix_rows,
    )

    rollup = json.loads((KEY_DATA / "scripturevec_key_results_rollup.json").read_text())
    rollup.update(
        {
            "created_from": str(EXPANDED_SUMMARY.relative_to(ROOT)),
            "layer_localization_completed_cells": int(data["result_file_count"]),
            "layer_localization_expected_cells": int(data["result_file_count"]),
            "layer_localization_missing_cells": [],
            "layer_localization_rows": int(data["row_count"]),
            "layer_localization_paired_rescue_count": int(data["paired_rescue_count"]),
            "top_layer_alpha_cells": cell_rows[:12],
            "chapter_stability": target_rows,
        }
    )
    (KEY_DATA / "scripturevec_key_results_rollup.json").write_text(json.dumps(rollup, indent=2) + "\n")

    readme = f"""# ScriptureVec Paper Key Data Guide

This folder contains compact, AI-readable data extracts for drafting the ScriptureVec Justice paper. The files are derived from local experiment artifacts under `results/experiments/scripturevec14`.

## Topline Counts

- book_discovery_l10_candidate_count: `19`
- book_confirmation_l40_candidate_count: `19`
- book_confirmation_l40_survivor_count: `7`
- chapter_discovery_l10_clean_hit_count: `38`
- chapter_confirmation_l40_candidate_count: `38`
- chapter_confirmation_l40_survivor_count: `16`
- layer_localization_completed_cells: `{data["result_file_count"]}`
- layer_localization_expected_cells: `{data["result_file_count"]}`
- layer_localization_rows: `{data["row_count"]}`
- layer_localization_paired_rescue_count: `{data["paired_rescue_count"]}`

## Files

- `book_discovery_l10_candidates.csv`: The 19 preliminary book-level Justice candidates from the all-66-book discovery atlas.
- `book_confirmation_l40_all_candidates.csv`: All 19 preliminary book candidates retested at limit 40, including survivors and failures.
- `book_confirmation_l40_survivors.csv`: The seven book-level candidates that survived limit-40 confirmation.
- `chapter_discovery_l10_clean_hits.csv`: The 38 clean preliminary chapter hits from the seven confirmed book sources.
- `chapter_confirmation_l40_all_candidates.csv`: All 38 chapter candidates retested at limit 40.
- `chapter_confirmation_l40_survivors.csv`: The 16 confirmed chapter movers used for layer localization.
- `layer_alpha_cells.csv`: Completed layer/alpha localization cells, ranked by rescue count.
- `layer_alpha_expected_grid.csv`: The completed expanded localization grid in planned order.
- `chapter_stability_by_localization.csv`: How often each confirmed chapter was rescued across completed localization cells.
- `layer_alpha_target_rows.csv`: One row per chapter per completed layer/alpha cell.
- `chapter_x_layer_alpha_rescue_matrix.csv`: Plot-ready binary matrix for chapter-by-cell rescue heatmaps.
- `scripturevec_key_results_rollup.json`: One-file structured summary of the packet.

## Important Note

The expanded localization grid is complete. The strongest broad-coverage cell is `L30 / alpha 96`, with 13 paired rescues out of 16 confirmed chapter vectors and mean positive delta 0.0875. The strongest high-movement cell remains `L24 / alpha 32`, with 9 paired rescues and mean positive delta 0.1687. The most efficient low-alpha cell is `L28 / alpha 16`, with 11 paired rescues at a much gentler steering strength. This supports the paper's central claim that specific biblical chapter vectors form a structured layer-and-strength map rather than a flat Scripture effect.
"""
    (KEY_DATA / "README.md").write_text(readme)


def font(size: int, bold: bool = False, italic: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_REG
    if bold:
        path = FONT_BOLD
    if italic:
        path = FONT_ITALIC
    return ImageFont.truetype(path, size=size)


def text_size(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur: list[str] = []
    for word in words:
        test = " ".join(cur + [word])
        if text_size(draw, test, fnt)[0] <= max_width or not cur:
            cur.append(word)
        else:
            lines.append(" ".join(cur))
            cur = [word]
    if cur:
        lines.append(" ".join(cur))
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    fnt: ImageFont.FreeTypeFont,
    fill: str,
    max_width: int,
    line_gap: int = 8,
    anchor: str = "la",
) -> int:
    x, y = xy
    lines = wrap_text(draw, text, fnt, max_width)
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill, anchor=anchor)
        y += fnt.size + line_gap
    return y


def clean_md_inline(text: str) -> str:
    text = text.replace("—", "—")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def update_markdown(md: str) -> str:
    md = md.replace(
        "**Author:** [Author Name], Institute for a Christian Machine Intelligence",
        "**Author:** Lucius, Institute for a Christian Machine Intelligence",
    )
    md = md.replace(
        "**Code & Data:** [repo link / DOI]",
        "**Code & Data:** https://github.com/christian-machine-intelligence/scripture-vector-steer",
    )
    md = re.sub(
        r"\n\[\*\*Figure 1 placeholder\.\*\* Discovery funnel.*?\]\n",
        "\n",
        md,
        flags=re.S,
    )
    md = md.replace("[**Figure 2 placeholder.** Book-level confirmation", "[**Figure 1 placeholder.** Book-level confirmation")
    md = md.replace("[**Figure 3 placeholder.** Chapter-discovery hits", "[**Figure 2 placeholder.** Chapter-discovery hits")
    md = re.sub(
        r"\n\[\*\*Table 2 placeholder\.\*\* The sixteen confirmed chapter movers.*?\]\n",
        "\nA compact table of the confirmed chapters appears in §6, after the localization results, where chapter identity and chapter behavior can be read together.\n",
        md,
        flags=re.S,
    )
    md = md.replace(
        "layer 31 at α = 96 rescued the broadest set of chapters (12/16), "
        "layer 32 at α = 32 rescued nearly as many (11/16) at a third the steering strength, "
        "and layer 24 at α = 32 produced the strongest mean positive movement on a smaller chapter set (9/16).",
        "an expanded 43-cell layer-and-strength grid found three distinct regimes: layer 30 at α = 96 rescued the broadest set of chapters (13/16), layer 28 at α = 16 rescued eleven chapters at a much gentler steering strength, and layer 24 at α = 32 produced the strongest mean positive movement on a smaller chapter set (9/16).",
    )
    md = md.replace(
        "A focused layer-and-strength localization grid then showed that these chapter vectors produced structured uptake: an expanded 43-cell layer-and-strength grid found three distinct regimes:",
        "An expanded 43-cell layer-and-strength localization grid then showed that these chapter vectors produced structured uptake in three distinct regimes:",
    )
    md = md.replace(
        "with broadest coverage at layer 31 (α = 96) and strongest mean movement at layer 24 (α = 32).",
        "with broadest coverage at layer 30 (α = 96), broad low-alpha uptake at layer 28 (α = 16), and strongest mean movement at layer 24 (α = 32).",
    )
    md = md.replace(
        "Discovery funnel from sixty-six biblical books to twenty-two completed layer/α localization cells.",
        "Discovery funnel from sixty-six biblical books to forty-three completed layer/α localization cells.",
    )
    md = md.replace(
        "The runtime-α ladder for the localization study was {16, 24, 32, 48, 64, 96}, paired with center layers {24, 28, 29, 30, 31, 32, 33, 36} in a focused grid: every center received the α = 32 condition, and centers 29, 31, and 32 received the full ladder.",
        "The runtime-α ladder for the localization study was {16, 24, 32, 48, 64, 96}, paired with center layers {24, 28, 29, 30, 31, 32, 33, 36} in an expanded focused grid. The final packet includes full alpha ladders for layers 24, 28, 29, 30, 31, 32, and 33, plus a layer-36 anchor cell at α = 32. Layer 36 was not expanded after its anchor cell produced no rescues.",
    )
    md = md.replace(
        "localized chapter-by-layer cells",
        "localized chapter-by-layer-and-strength cells",
    )
    md = md.replace(
        "The localization study tested the sixteen confirmed chapter vectors across a focused layer-by-α grid. "
        "The full design specified twenty-three behavior cells; twenty-two were completed in the present packet. "
        "A single cell — center layer 31, α = 48 — remained as a backfill that did not launch before the data freeze, "
        "and we report the twenty-two completed cells with the missing cell flagged as missing data.",
        "The localization study tested the sixteen confirmed chapter vectors across an expanded focused layer-by-α grid. "
        "The final packet contains forty-three completed behavior cells: full α ladders for layers 24, 28, 29, 30, 31, 32, and 33, plus the layer-36 anchor cell at α = 32.",
    )
    md = md.replace(
        "Layer-by-α rescue heatmap, x-axis α, y-axis center layer, cell value paired-rescue count out of sixteen. "
        "The L31/α 48 cell is marked as a missing backfill.",
        "Layer-by-α rescue heatmap, x-axis α, y-axis center layer, cell value paired-rescue count out of sixteen. "
        "The expanded 43-cell grid is complete; unexpanded layer-36 cells are marked as not run.",
    )
    new_section_5 = """## 5. Results: Layer and α Localization

### 5.1 The Expanded Localization Grid

The localization study tested the sixteen confirmed chapter vectors across an expanded focused layer-by-α grid. The final packet contains forty-three completed behavior cells: full α ladders for layers 24, 28, 29, 30, 31, 32, and 33, plus the layer-36 anchor cell at α = 32. The layer-36 anchor produced no rescues, so it was intentionally not expanded.

[**Figure 3 placeholder.** Layer-by-α rescue heatmap, x-axis α, y-axis center layer, cell value paired-rescue count out of sixteen. The expanded 43-cell grid is complete; unexpanded layer-36 cells are marked as not run. Generated from `layer_alpha_expected_grid.csv`.]

### 5.2 Three Regimes of Biblical Justice Steering

The expanded grid revealed three distinct regimes rather than a single best setting. The broadest rescue regime was layer 30 at α = 96, which rescued thirteen of the sixteen confirmed chapter vectors at a mean positive Δ of 0.0875. The neighboring layer 31 at α = 96 cell rescued twelve of sixteen at mean Δ = 0.0812, confirming that broad coverage lives in the late-middle high-alpha region rather than in a lone accident.

The efficient-uptake regime was layer 28 at α = 16. That cell rescued eleven of sixteen chapters at mean Δ = 0.0625 while using the gentlest steering strength in the tested ladder. Its importance is not that it moved the model the most, but that many chapter-derived vectors became useful without brute-force intervention.

The high-force selective regime was layer 24 at α = 32. It rescued only nine of sixteen chapters, but it produced the strongest mean positive movement in the grid, Δ = 0.1687. Lower-layer steering was therefore powerful but thresholded: at α = 32 it produced large gains, while at α = 48, 64, and 96 it collapsed to zero paired rescues and negative mean movement.

### 5.3 Breadth, Strength, and Thresholds

Plotting paired rescue count against mean positive Δ across the forty-three completed cells exposed a tradeoff that the rescue counts alone obscured. L30/α96 maximized breadth; L24/α32 maximized strength; L28/α16 established that broad uptake could appear at low steering strength. These are different kinds of success. A single undifferentiated Justice direction would be expected to peak in breadth and strength together. Instead, the observed peaks were displaced across layer and α.

[**Figure 4 placeholder.** Breadth vs strength scatter: x = paired rescue count, y = mean positive Δ, points labeled with `L{layer}/α{α}`, the three regime-defining cells highlighted. Generated from `layer_alpha_cells.csv`.]

The α trajectories make the threshold pattern visible. Layer 24 rises sharply at α = 32 and then collapses under heavier pushes. Layer 28 is best at α = 16 and weakens as α rises. Layer 30, by contrast, improves with stronger pushes and reaches the broadest coverage at α = 96. The model is therefore not merely responding to "more vector"; different layers admit different amounts and kinds of biblical-justice signal.

[**Figure 5 placeholder.** Alpha trajectories by layer: x = α, y = paired-rescue count out of sixteen, with layers 24, 28, 30, 31, 32, and 33 shown as separate trajectories. Generated from `layer_alpha_cells.csv`.]

This tradeoff was the most mechanistically suggestive result in the paper. The chapter-derived Justice direction appears internally compound: some components become useful through gentle lower-middle-layer uptake, some through high-alpha late-middle steering, and some through a narrow lower-layer threshold. Activation steering can map that structure behaviorally; sparse-feature work can later test which internal features mediate it.

"""
    md = re.sub(
        r"## 5\. Results: Layer and α Localization\n.*?(?=## 6\. Chapter Stability and Differential Uptake)",
        new_section_5,
        md,
        flags=re.S,
    )
    new_section_6 = """## 6. Chapter Stability and Differential Uptake

The localization grid afforded a second, orthogonal view of the chapter set. Across the forty-three completed cells, some chapters were rescued in many cells and others in only a few. Acts 11 was rescued in twenty-nine cells; Acts 7 in twenty; 1 Chronicles 9 and Acts 16 in nineteen each; Acts 27 in eighteen; Hebrews 2 in seventeen; and Numbers 27 in sixteen. Numbers 22, by contrast, surfaced in eight cells. The chapter set therefore has a stable core and a more selective edge.

[**Figure 6 placeholder.** Chapter stability across layer-α cells, ranked horizontal bar chart. Generated from `chapter_stability_by_localization.csv`.]

At this point, the sixteen confirmed chapters can be read in two ways at once: first as confirmed chapter-level movers from the limit-40 run, and second as differently stable vectors across the localization grid. Table 2 is therefore not just a list of discoveries; it is the bridge between the confirmation result and the layer/strength result.

[**Table 2 placeholder.** The sixteen confirmed chapter movers, with biblical reference, book family, limit-40 control accuracy, limit-40 positive-steering accuracy, negative-α and null accuracies, localization rescue count, best layer/α cell, and brief biblical motif. Generated from `chapter_confirmation_l40_survivors.csv` and `chapter_stability_by_localization.csv`.]

The structure grew more interesting still when the chapter-by-cell rescue matrix was read across α at fixed layer. At layer 28, α = 16 surfaced a broad New Testament-heavy family, especially Acts and Hebrews. At layer 30, the rescued family widened as α rose, reaching thirteen chapters at α = 96 and bringing in Deuteronomy, Judges, Numbers, Acts, Hebrews, and Chronicles together. At layer 24, by contrast, α = 32 concentrated high-movement effects in a smaller cluster, while heavier α values collapsed. Steering strength therefore selected different components of the chapter-derived geometry rather than simply amplifying the same effect at higher volume.

[**Figure 7 placeholder.** Chapter-by-layer/α rescue matrix, sixteen rows by forty-three completed columns, binary heatmap. Generated from `chapter_x_layer_alpha_rescue_matrix.csv`.]

The Amos result discussed in §4.4 belongs here as well. Amos confirmed at the book level — its prophetic call for justice to *"run down as waters"* (Amos 5:24, KJV) is among the most justice-saturated rhetoric in Scripture — and yet no Amos chapter survived the limit-40 chapter confirmation. The pattern is consistent with a book-level vector aggregating signal distributed across the whole book: the prophet's sustained indictment of unjust commerce, false worship, and forgotten orphans runs from chapter to chapter rather than concentrating in any one of them. §7 returns to the point.

"""
    md = re.sub(
        r"## 6\. Chapter Stability and Differential Uptake\n.*?(?=## 7\. Biblical Patterning)",
        new_section_6,
        md,
        flags=re.S,
    )
    md = re.sub(
        r"\n\*\*Missing localization cell\.\*\*.*?(?=\n\n\*\*Translation dependence\.\*\*)",
        "\n",
        md,
        flags=re.S,
    )
    md = re.sub(
        r"\n\*\*Backfill the missing localization cell\.\*\*.*?(?=\n\n\*\*Cross-model replication\.\*\*)",
        "\n",
        md,
        flags=re.S,
    )
    md = md.replace(
        "the localization study then placed those sixteen vectors on a layer-by-α map whose three flagship cells distributed the work — L31/α 96 for breadth, L32/α 32 for broad coverage at lower strength, L24/α 32 for the strongest per-chapter movement.",
        "the localization study then placed those sixteen vectors on a layer-by-α map whose three regimes distributed the work — L30/α96 for broadest rescue, L28/α16 for efficient low-strength uptake, and L24/α32 for the strongest per-chapter movement.",
    )
    md = md.replace(
        "Original experiment summaries (`results/experiments/scripturevec14/`); the compact data packet used for figures and tables (`key_data/`); design documents for the discovery, confirmation, and localization runs; and the scripts used for compact-data summarization. The missing cell — layer 31 at α = 48 — is documented as a backfill that did not launch before the data freeze and should be backfilled before final publication.",
        "Original experiment summaries (`results/experiments/scripturevec14/`); the compact data packet used for figures and tables (`key_data/`); design documents for the discovery, confirmation, and localization runs; and the scripts used for compact-data summarization. The expanded 43-cell localization grid is complete and included in the current summary packet.",
    )
    md = md.replace(
        "the binary chapter-by-cell rescue matrix used to generate Figure 7.",
        "the binary chapter-by-cell rescue matrix used to generate Figure 7.",
    )
    md = md.replace(
        "The picture this work presents is of scripture-as-structured-geometry.",
        "The picture this work presents is of scripture-as-structured-geometry: Scripture can be treated not merely as alignment text, but as a structured source of discoverable moral steering vectors whose behavioral force can be mapped across a model's internal geometry.",
    )
    md = md.replace("the L24-vs-L31/L32 split", "the L24-vs-L28/L30 split")
    md = md.replace("the L24 and L31/L32 effects", "the L24, L28, and L30 effects")
    md = md.replace(
        "The breadth-versus-strength split (§5.3) and the chapter-by-α selectivity (§6) carried the paper's most mechanistically suggestive load. The straightforward reading is that the chapter-derived Justice direction is internally compound: the steering effect at L24 and the steering effect at L31/L32 picked up different sub-structures of that compound representation. The chapter-by-α matrix in Figure 7 stands as the paper's strongest invitation to sparse-feature analysis, since rows of the matrix that activated together under the same α constitute a behavioral fingerprint that should map onto a small number of sparse features in a trained sparse autoencoder. The L24-vs-L31/L32 split sharpens the same point: the right granularity for explanation is the Justice *complex*, a small number of distinguishable directions whose superposition produced the observed behavior. Activation steering can motivate these claims; only sparse-feature work can settle them.",
        "The breadth-versus-strength split (§5.3), the alpha trajectories in Figure 5, and the chapter-by-α selectivity (§6) carried the paper's most mechanistically suggestive load. The straightforward reading is that the chapter-derived Justice direction is internally compound: the L24 threshold effect, the L28 low-alpha effect, and the L30 high-alpha broad-rescue effect pick up different sub-structures of that compound representation. The chapter-by-α matrix in Figure 7 stands as the paper's strongest invitation to sparse-feature analysis, since rows of the matrix that activated together under the same α constitute a behavioral fingerprint that should map onto a small number of sparse features in a trained sparse autoencoder. The right granularity for explanation is the Justice *complex*, a small number of distinguishable directions whose superposition produced the observed behavior. Activation steering can motivate these claims; only sparse-feature work can settle them.",
    )
    md = md.replace(
        'The third was the inclusion of Hebrews 7. A modern reader looking for a justice text would not, on first instinct, reach for a chapter on Melchizedekian priesthood; biblically, however, *"king of righteousness"* and *"king of peace"* is justice in its enduring priestly mode, and the chapter rose in coverage at the higher-α cells that surfaced the Hebrews family. The strangeness of the Hebrews 7 hit is part of the result: the model proved sensitive to biblical scenes of order, adjudication, mediation, and recompense rather than to the surface vocabulary of modern justice.',
        'The third was the inclusion of Hebrews 7. Its surface topic is Melchizedekian priesthood, yet *"king of righteousness"* and *"king of peace"* is justice in its enduring priestly mode. The chapter was strongest in the L24/α32 selective regime, which makes it a useful test case for whether the model is responding to biblical forms of order, mediation, and right office rather than to justice vocabulary alone.',
    )
    md = md.replace(
        "**Discovery screens are not confirmation screens.** Limit-10 screens function as behavioral discovery: they exist to filter, not to confirm. The seven confirmed book sources and the sixteen confirmed chapter movers, by contrast, passed limit-40 confirmation. The paper's load-bearing claims are at the limit-40 level, and readers should keep this distinction in mind.",
        "**Discovery and localization screens are not confirmation screens.** Limit-10 screens function as behavioral discovery: they exist to filter and localize, not to confirm at final scale. The seven confirmed book sources and the sixteen confirmed chapter movers passed limit-40 confirmation; the expanded layer/α atlas remains limit-10 localization evidence. The next confirmation step is to retest the regime-defining cells on a wider slice.",
    )
    further_work_insert = (
        "\n**Confirm the regime-defining localization cells.** The expanded grid identifies L30/α96, L28/α16, and L24/α32 as the most important regimes. "
        "A wider-slice confirmation should retest these settings, and likely L31/α96 as the neighboring broad-coverage comparison, before the strongest localization claims are treated as final.\n"
    )
    md = md.replace(
        "\n**Cross-model replication.** Replicating the canon-wide pipeline",
        further_work_insert + "\n**Cross-model replication.** Replicating the canon-wide pipeline",
    )
    new_section_7 = """### 7.2 Why These Chapters Moved the Justice Benchmark

The commentary tradition matters here because it helps name the kind of justice the benchmark appears to be receiving from these chapter vectors. The ratio stage of VirtueBench2 asks the model to hold the just answer when a plausible rationale for the unjust answer is placed under its nose. The confirmed chapters are dense with the biblical forms that train exactly that posture: judgment rendered under God, claims heard and answered, corrupt incentives refused, offices restored to their right order, testimony preserved under pressure, and recompense returned to the one to whom it is due.

The Pentateuchal chapters show this in legal and quasi-legal form. Deuteronomy 16 does not merely contain the word justice; it binds public worship to incorruptible judgment, appointing judges and forbidding bribes before the command to follow what is altogether just. Numbers 27 stages a claim, a hearing, a divine verdict, and an amended inheritance rule. These chapters plausibly help because their vectors carry justice as ordered adjudication rather than as sentiment. Numbers 11 and Numbers 22 add two pressure cases. In Numbers 11, disordered appetite is judged and Moses' burden is distributed through appointed elders. In Numbers 22, speech offered for hire is constrained by God. Both map naturally onto a benchmark setting where the model must resist a tempting but wrong justification.

The historical chapters add the political and institutional face of the same pattern. Judges 7 makes deliverance impossible to misattribute; the victory is ordered so that Israel cannot claim for itself what belongs to the Lord. Judges 9 is a severe narrative of usurpation and recompense, closing with wickedness rendered back upon Abimelech and Shechem. 1 Chronicles 9 and 29 are quieter but not weaker: restored offices, ordered worship, and David's confession that Israel gives only what has first come from God. Those are not abstract moral labels. They are forms of right relation - right office, right attribution, right stewardship - and a benchmark about justice can receive them as pressure against self-serving answers.

Acts supplies the clearest public-testimony cluster. Stephen's speech in Acts 7 names Christ as the Just One after rehearsing a history of rejected deliverers. Acts 11 turns Peter's testimony into an ecclesial verdict about Gentile inclusion and then into relief sent according to ability. Acts 16 exposes unlawful punishment and insists on public accountability from magistrates. Acts 27 places truthful counsel in the mouth of the prisoner whom the authorities should have heeded. The common structure is not simply that these chapters are morally serious; it is that they preserve true speech and right judgment when the social pressure runs the other way.

Hebrews 2, 9, and 10 supply the priestly and eschatological side of the same justice grammar. Hebrews 2 names just recompense directly and ties it to the faithful high priest who shares the condition of those he redeems. Hebrews 9 joins sacrifice, conscience, inheritance, and final judgment. Hebrews 10 binds vengeance, recompense, and the life of the just under faith. These chapters make the benchmark-relevant pressure more ultimate: justice is not only human procedure, but a divine ordering in which wrong, mediation, and judgment cannot be separated.

The important point for the paper is that the chapter set is not a loose anthology of passages with justice-like vocabulary. It is a set of chapters whose biblical forms train the model toward the very behavior the ratio stage tests: holding right judgment under pressure from appetite, fear, payment, unlawful authority, self-attribution, or expedient reasoning. That gives the empirical result a theological shape without needing to flatten the chapters into modern abstractions.

### 7.3 Hebrews 7 as a Special Case

Hebrews 7 deserves separate treatment because it is one of the most interesting discoveries in the set. It survived the limit-40 chapter confirmation and was rescued in eleven of the forty-three localization cells, with its strongest observed cell at L24/α32 - the same lower-layer, high-movement regime that produced the largest mean positive shift in the expanded grid. It is not the broadest chapter in the atlas, but it is a sharp one.

At first glance, Hebrews 7 may look like a strange Justice hit. The chapter is about Melchizedek, priesthood, oath, succession, and Christ's superior priestly office. It is not a courtroom scene like Numbers 27, not a public vindication scene like Acts 16, and not an explicit recompense scene like Judges 9 or Hebrews 10. That is exactly why the result is valuable. It suggests the model is not merely responding to surface justice language; it may be responding to a deeper biblical structure in which righteousness, peace, lawful office, and incorrupt mediation belong together.

The chapter itself makes that structure explicit through Melchizedek's name and title: king of righteousness and king of peace. Its argument turns on the insufficiency of a merely inherited office and the emergence of a priesthood grounded in oath, permanence, holiness, and indestructible life. In the logic of Hebrews, justice is not only the rendering of verdicts; it is the establishment of a mediator who can actually put persons, covenant, and God in right relation. That gives the Hebrews 7 vector a different texture from the Acts or Numbers vectors.

This matters for interpreting the steering result. VirtueBench2's Justice ratio items ask the model to reject answers that can be made to sound prudent, useful, or convenient while still violating what is due. A Hebrews 7 direction may help by activating a representation of justice as right mediation and rightful office: the answer must be ordered by what is true and due, not merely by what seems expedient. That would explain why Hebrews 7 becomes especially useful at the L24/α32 selective regime. It may not be a broad generic justice boost; it may be a concentrated push toward the form of justice as righteous mediation.

For the paper's argument, Hebrews 7 is therefore not an awkward exception. It is a proof that the pipeline can surface chapters whose relevance becomes clear only when biblical justice is read in its own register. The result asks for commentary-backed interpretation because the chapter's justice content is priestly, typological, and institutional rather than procedural on the surface.

"""
    md = re.sub(
        r"### 7\.2 The Confirmed Chapters Against the Commentary Tradition\n.*?(?=### 7\.3 Four Theses About the Collection)",
        new_section_7,
        md,
        flags=re.S,
    )
    md = md.replace("### 7.3 Four Theses About the Collection", "### 7.4 Four Theses About the Collection")
    md = md.replace("### 7.4 What the Chapter Set Shows", "### 7.5 What the Chapter Set Shows")
    md = md.replace("commentary readings of §7.2", "commentary readings of §§7.2-7.3")
    md = re.sub(
        r"\n\[\*\*Table 3 placeholder\.\*\* Biblical patterning.*?\]\n",
        "\n",
        md,
        flags=re.S,
    )
    md = re.sub(
        r"### 7\.5 What the Chapter Set Shows\n.*?(?=---\n\n## 8\. Discussion)",
        "",
        md,
        flags=re.S,
    )
    revised_82 = """### 8.2 Anomalies and Tensions

Three results invited careful theological reading. The first was the Amos asymmetry. Amos is the most prophetically justice-rhetorical book in the confirmed set, yet no individual Amos chapter survived chapter confirmation. That does not weaken the book-level result. It sharpens it. The book vector appears to have recovered a distributed prophetic argument that was not compact enough to be captured by any single chapter vector.

This resembles the compactness pattern reported in ICMI-010, §4.3. There, the decisive moral movement did not live in a bare proof-text alone: James 4:17 by itself was not sufficient, and the surrounding scriptural framework without the verse was also not sufficient. The effect appeared in the assembled moral unit. Amos suggests the same principle at the book scale. The model may recognize Amos as a coherent prophetic indictment — false worship, unjust commerce, oppression of the poor, and divine judgment distributed across the whole book — while no one chapter carries enough of that shape to survive as an independent steering direction.

That makes Amos one of the paper's strongest signs that there is a real aspect of Scripture inside the model rather than a mere keyword response. A keyword account would expect Amos 5, with its famous justice language, to dominate at chapter level. Instead, the book survives and the chapters fall away. The effect therefore looks less like lexical recognition and more like book-scale moral form: the model has received the book as a whole argument whose force is distributed across its parts.

The second tension was the L24-vs-L28/L30 split. The divergence of the layer producing the strongest per-chapter movement from the layer producing the broadest rescue is consistent with a chapter-derived direction that is internally compound, with different sub-features picked up at different layers and amplified at different strengths. The third was the inclusion of Hebrews 7. Its surface topic is Melchizedekian priesthood, yet *"king of righteousness"* and *"king of peace"* is justice in its enduring priestly mode. The chapter was strongest in the L24/α32 selective regime, which makes it a useful test case for whether the model is responding to biblical forms of order, mediation, and right office rather than to justice vocabulary alone.

### 8.3 Mechanistic Implications

"""
    md = re.sub(
        r"### 8\.2 Anomalies and Tensions\n.*?(?=The breadth-versus-strength split)",
        revised_82,
        md,
        flags=re.S,
    )
    md = md.replace("### 8.4 Mechanistic Implications", "### 8.3 Mechanistic Implications")
    revised_conclusion = """## 11. Conclusion

This paper began with a simple question: whether Scripture inside a language model can be searched, resolved, and reused as structured moral geometry. The answer returned by the experiment is yes, at least for the Justice slice of *VirtueBench2* in Qwen3-14B. The useful signal did not appear as an undifferentiated biblical atmosphere. It resolved to particular books, then to particular chapters, then to particular layer-and-strength regimes. That resolution is the main finding. Scripture can be treated not merely as alignment text placed before a model, but as a structured source of discoverable moral steering vectors whose behavioral force can be mapped across a model's internal geometry.

The confirmed chapters give that claim its theological specificity. Numbers 27, Acts 11, Hebrews 7, and the rest do not simply share a modern justice keyword. They perform recognizable biblical forms of justice: claim and adjudication, right office, faithful testimony under pressure, mediation, recompense, restored worship, and stewardship before God. When chapter-derived vectors from these texts moved the model toward the just answer, the result suggested that the model had received more than isolated phrases. It had internalized enough of Scripture's moral form that some of that form could be recovered as an intervention.

The work therefore changes the shape of the next question. The question is no longer only whether scriptural material can improve an alignment benchmark. It is whether the canon's internal moral structure can be discovered at useful resolution inside trained models, compared across models, and eventually decomposed into the sparse features that mediate its effects. Activation steering has shown the behavioral map. The next interpretability step is to read the map at finer resolution and ask which features carry recompense, adjudication, righteous mediation, and speech held steady under pressure.

*"The work of righteousness shall be peace; and the effect of righteousness, quietness and assurance for ever"* (Isaiah 32:17, KJV). The present study does not claim that a model possesses that righteousness. It claims something narrower and still remarkable: that Scripture's articulation of justice has left a structured, measurable trace in model space, and that this trace can be searched, steered, localized, and prepared for mechanistic explanation.

"""
    md = re.sub(
        r"## 11\. Conclusion\n.*?(?=---\n\n## Appendix A:)",
        revised_conclusion,
        md,
        flags=re.S,
    )
    if "ICMI-010." not in md:
        md = md.replace(
            "ICMI-020. (2026). *Beyond the Psalm: Canon-Wide Scripture Injection and Virtue-Evaluation Behavior*. ICMI Working Paper No. 020.",
            "ICMI-020. (2026). *Beyond the Psalm: Canon-Wide Scripture Injection and Virtue-Evaluation Behavior*. ICMI Working Paper No. 020.\n\nICMI-010. (2026). *Moral Compactness and Distributed Scriptural Force*. ICMI Working Paper No. 010. https://icmi-proceedings.com/ICMI-010-moral-compactness.html",
        )
    return md


def axis(draw: ImageDraw.ImageDraw, x0: int, y0: int, x1: int, y1: int) -> None:
    draw.line((x0, y0, x1, y0), fill="#94a3b8", width=2)
    draw.line((x0, y0, x0, y1), fill="#94a3b8", width=2)


def title(draw: ImageDraw.ImageDraw, text: str, subtitle: str, w: int) -> None:
    draw.text((70, 45), text, font=font(44, bold=True), fill=NAVY)
    if subtitle:
        draw.text((72, 103), subtitle, font=font(24), fill=GRAY)


def save_img(img: Image.Image, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG", optimize=True)


def figure_1() -> Path:
    rows = read_csv("book_confirmation_l40_all_candidates.csv")
    rows.sort(key=lambda r: (r["survived"] != "yes", -float(r["positive_delta"]), r["reference"]))
    w, h = 1800, 1180
    img = Image.new("RGB", (w, h), PAPER)
    d = ImageDraw.Draw(img)
    title(d, "Figure 1. Book Confirmation", "Limit-40 retest of 19 preliminary book candidates; blue rows are the seven survivors", w)
    left, right, top, bottom = 260, 1660, 185, 1060
    min_x, max_x = 0.34, 0.46
    def sx(v: float) -> int:
        return int(left + (v - min_x) / (max_x - min_x) * (right - left))
    for tick in [0.35, 0.375, 0.40, 0.425, 0.45]:
        x = sx(tick)
        d.line((x, top, x, bottom), fill="#e2e8f0", width=2)
        d.text((x, bottom + 18), f"{tick:.3f}".rstrip("0").rstrip("."), font=font(20), fill=GRAY, anchor="ma")
    d.line((sx(0.4), top, sx(0.4), bottom), fill="#334155", width=3)
    n = len(rows)
    row_h = (bottom - top) / n
    for i, r in enumerate(rows):
        y = int(top + row_h * (i + 0.5))
        survived = r["survived"] == "yes"
        color = BLUE if survived else "#94a3b8"
        label_color = NAVY if survived else GRAY
        d.text((left - 22, y), r["reference"], font=font(21, bold=survived), fill=label_color, anchor="rm")
        c = sx(float(r["control_accuracy"]))
        p = sx(float(r["positive_accuracy"]))
        d.line((c, y, p, y), fill=color, width=5 if survived else 3)
        d.ellipse((c - 7, y - 7, c + 7, y + 7), fill="white", outline="#475569", width=3)
        d.ellipse((p - 10, y - 10, p + 10, y + 10), fill=color, outline=color)
        if survived:
            d.text((p + 18, y), "+", font=font(22, bold=True), fill=color, anchor="lm")
    d.text((sx(0.4), top - 25), "control baseline", font=font(20, bold=True), fill="#334155", anchor="ma")
    out = FIG_DIR / "figure_1_book_confirmation.png"
    save_img(img, out)
    return out


def figure_2() -> Path:
    disc = read_csv("chapter_discovery_l10_clean_hits.csv")
    conf = read_csv("chapter_confirmation_l40_survivors.csv")
    books = ["Acts", "Hebrews", "Numbers", "1 Chronicles", "Judges", "Amos", "Deuteronomy"]
    disc_counts = {b: 0 for b in books}
    conf_counts = {b: 0 for b in books}
    for r in disc:
        disc_counts[r["book"]] = disc_counts.get(r["book"], 0) + 1
    for r in conf:
        conf_counts[r["book"]] = conf_counts.get(r["book"], 0) + 1
    w, h = 1600, 950
    img = Image.new("RGB", (w, h), PAPER)
    d = ImageDraw.Draw(img)
    title(d, "Figure 2. Chapter Hits by Confirmed Book", "Preliminary chapter hits and limit-40 survivors by source book", w)
    left, right, top, bottom = 180, 1490, 190, 790
    max_v = max(disc_counts.values())
    for tick in range(0, max_v + 2, 2):
        y = int(bottom - tick / (max_v + 1) * (bottom - top))
        d.line((left, y, right, y), fill="#e2e8f0", width=2)
        d.text((left - 20, y), str(tick), font=font(20), fill=GRAY, anchor="rm")
    group_w = (right - left) / len(books)
    for i, b in enumerate(books):
        x0 = int(left + i * group_w + 28)
        base = bottom
        bw = 46
        for j, (value, color, label) in enumerate([(disc_counts[b], "#93c5fd", "prelim"), (conf_counts[b], BLUE, "confirmed")]):
            x = x0 + j * (bw + 14)
            y = int(bottom - value / (max_v + 1) * (bottom - top))
            d.rounded_rectangle((x, y, x + bw, base), radius=8, fill=color)
            d.text((x + bw / 2, y - 16), str(value), font=font(21, bold=True), fill=NAVY, anchor="ma")
        lines = wrap_text(d, b, font(21, bold=True), int(group_w - 18))
        y_label = bottom + 35
        for line in lines:
            d.text((x0 + 50, y_label), line, font=font(21, bold=True), fill=NAVY, anchor="ma")
            y_label += 25
    d.rounded_rectangle((1160, 130, 1490, 220), radius=12, fill="white", outline="#cbd5e1")
    d.rectangle((1190, 158, 1225, 188), fill="#93c5fd")
    d.text((1240, 173), "preliminary hits", font=font(21), fill=NAVY, anchor="lm")
    d.rectangle((1190, 193, 1225, 223), fill=BLUE)
    d.text((1240, 208), "confirmed movers", font=font(21), fill=NAVY, anchor="lm")
    out = FIG_DIR / "figure_2_chapter_hits_by_book.png"
    save_img(img, out)
    return out


def heat_color(v: int, max_v: int = 13) -> str:
    if v <= 0:
        return "#f8fafc"
    # Blend from pale blue to navy.
    t = v / max_v
    r1, g1, b1 = (219, 234, 254)
    r2, g2, b2 = (30, 64, 175)
    r = int(r1 + (r2 - r1) * t)
    g = int(g1 + (g2 - g1) * t)
    b = int(b1 + (b2 - b1) * t)
    return f"#{r:02x}{g:02x}{b:02x}"


def figure_3() -> Path:
    rows = read_csv("layer_alpha_expected_grid.csv")
    vals = {(int(r["center_layer"]), int(float(r["alpha"]))): int(r["paired_rescues"]) for r in rows}
    centers = [24, 28, 29, 30, 31, 32, 33, 36]
    alphas = [16, 24, 32, 48, 64, 96]
    w, h = 1400, 950
    img = Image.new("RGB", (w, h), PAPER)
    d = ImageDraw.Draw(img)
    title(d, "Figure 3. Layer/α Rescue Heatmap", "Cell value is paired rescues out of 16 confirmed chapter vectors", w)
    left, top = 250, 245
    cell_w, cell_h = 145, 78
    d.text((left + len(alphas) * cell_w / 2, top - 84), "Runtime α", font=font(25, bold=True), fill=NAVY, anchor="ma")
    for j, a in enumerate(alphas):
        d.text((left + j * cell_w + cell_w / 2, top - 25), str(a), font=font(24, bold=True), fill=NAVY, anchor="ma")
    d.text((115, top + len(centers) * cell_h / 2), "Center layer", font=font(25, bold=True), fill=NAVY, anchor="mm")
    for i, c in enumerate(centers):
        y = top + i * cell_h
        d.text((left - 35, y + cell_h / 2), str(c), font=font(24, bold=True), fill=NAVY, anchor="rm")
        for j, a in enumerate(alphas):
            x = left + j * cell_w
            if (c, a) in vals:
                v = vals[(c, a)]
                fill = heat_color(v)
                d.rounded_rectangle((x + 4, y + 4, x + cell_w - 4, y + cell_h - 4), radius=10, fill=fill, outline="#cbd5e1")
                txt = f"{v}/16"
                txt_color = "white" if v >= 9 else NAVY
                d.text((x + cell_w / 2, y + cell_h / 2), txt, font=font(27, bold=True), fill=txt_color, anchor="mm")
                if (c, a) in [(30, 96), (31, 96), (28, 16), (24, 32)]:
                    d.rounded_rectangle((x + 1, y + 1, x + cell_w - 1, y + cell_h - 1), radius=12, outline=AMBER, width=5)
            else:
                d.rounded_rectangle((x + 4, y + 4, x + cell_w - 4, y + cell_h - 4), radius=10, fill="#f1f5f9", outline="#e2e8f0")
                d.text((x + cell_w / 2, y + cell_h / 2), "not run", font=font(18), fill="#94a3b8", anchor="mm")
    out = FIG_DIR / "figure_3_layer_alpha_heatmap.png"
    save_img(img, out)
    return out


def figure_4() -> Path:
    rows = read_csv("layer_alpha_cells.csv")
    w, h = 1500, 980
    img = Image.new("RGB", (w, h), PAPER)
    d = ImageDraw.Draw(img)
    title(d, "Figure 4. Breadth Versus Strength", "Breadth is paired-rescue count; strength is mean positive delta within the cell", w)
    left, right, top, bottom = 170, 1360, 175, 790
    max_x = 13
    max_y = 0.18
    for xval in range(0, 14, 2):
        x = int(left + xval / max_x * (right - left))
        d.line((x, top, x, bottom), fill="#e2e8f0", width=2)
        d.text((x, bottom + 18), str(xval), font=font(20), fill=GRAY, anchor="ma")
    for yval in [0, 0.04, 0.08, 0.12, 0.16]:
        y = int(bottom - yval / max_y * (bottom - top))
        d.line((left, y, right, y), fill="#e2e8f0", width=2)
        d.text((left - 18, y), f"{yval:.2f}", font=font(20), fill=GRAY, anchor="rm")
    axis(d, left, bottom, right, top)
    d.text(((left + right) / 2, bottom + 65), "Paired rescues out of 16", font=font(24, bold=True), fill=NAVY, anchor="ma")
    d.text((left, top - 34), "Mean positive delta", font=font(24, bold=True), fill=NAVY, anchor="la")
    highlights = {(30, 96), (31, 96), (28, 16), (24, 32)}
    label_offsets = {
        (30, 96): (-165, -34),
        (31, 96): (-150, 18),
        (28, 16): (20, 18),
        (24, 32): (20, -34),
    }
    for r in rows[::-1]:
        c = int(r["center_layer"])
        a = int(float(r["alpha"]))
        xval = int(r["paired_rescues"])
        yval = float(r["mean_positive_delta"])
        x = int(left + xval / max_x * (right - left))
        y = int(bottom - yval / max_y * (bottom - top))
        high = (c, a) in highlights
        color = AMBER if high else (BLUE if c in [31, 32] else "#94a3b8")
        rad = 14 if high else 9
        d.ellipse((x - rad, y - rad, x + rad, y + rad), fill=color, outline="white", width=3)
        if high:
            label = f"L{c}/α{a}"
            dx, dy = label_offsets[(c, a)]
            d.text((x + dx, y + dy), label, font=font(22, bold=True), fill=NAVY, anchor="la")
    out = FIG_DIR / "figure_4_breadth_strength.png"
    save_img(img, out)
    return out


def figure_5() -> Path:
    rows = read_csv("layer_alpha_cells.csv")
    values = {
        (int(r["center_layer"]), int(float(r["alpha"]))): int(r["paired_rescues"])
        for r in rows
    }
    layers = [24, 28, 30, 31, 32, 33]
    alphas = [16, 24, 32, 48, 64, 96]
    colors = {
        24: AMBER,
        28: TEAL,
        30: BLUE,
        31: NAVY,
        32: GREEN,
        33: "#7c3aed",
    }
    w, h = 1650, 980
    img = Image.new("RGB", (w, h), PAPER)
    d = ImageDraw.Draw(img)
    title(d, "Figure 5. Alpha Trajectories by Layer", "Different layers respond to steering strength in different ways", w)
    left, right, top, bottom = 190, 1375, 190, 790
    max_y = 13
    def sx(alpha: int) -> int:
        return int(left + (alphas.index(alpha) / (len(alphas) - 1)) * (right - left))
    def sy(v: int) -> int:
        return int(bottom - v / max_y * (bottom - top))
    for tick in [0, 3, 6, 9, 12, 13]:
        y = sy(tick)
        d.line((left, y, right, y), fill="#e2e8f0", width=2)
        d.text((left - 18, y), str(tick), font=font(20), fill=GRAY, anchor="rm")
    for a in alphas:
        x = sx(a)
        d.line((x, top, x, bottom), fill="#eef2f7", width=2)
        d.text((x, bottom + 24), str(a), font=font(21, bold=True), fill=NAVY, anchor="ma")
    axis(d, left, bottom, right, top)
    d.text(((left + right) / 2, bottom + 72), "Runtime α", font=font(24, bold=True), fill=NAVY, anchor="ma")
    d.text((left, top - 36), "Paired rescues out of 16", font=font(24, bold=True), fill=NAVY, anchor="la")
    for layer in layers:
        pts = [(sx(a), sy(values.get((layer, a), 0))) for a in alphas]
        color = colors[layer]
        d.line(pts, fill=color, width=5)
        for x, y in pts:
            d.ellipse((x - 8, y - 8, x + 8, y + 8), fill=color, outline="white", width=3)
        lx, ly = pts[-1]
        d.text((lx + 16, ly), f"L{layer}", font=font(22, bold=True), fill=color, anchor="lm")
    note = "L24 spikes then collapses; L28 is strongest at gentle α; L30 rises into the broadest high-α rescue."
    d.rounded_rectangle((185, 845, 1465, 920), radius=16, fill="#eef6ff", outline="#bfdbfe", width=2)
    d.text((220, 882), note, font=font(24, bold=True), fill=NAVY, anchor="lm")
    out = FIG_DIR / "figure_5_alpha_trajectories.png"
    save_img(img, out)
    return out


BOOK_COLORS = {
    "Acts": BLUE,
    "Hebrews": TEAL,
    "Numbers": AMBER,
    "1 Chronicles": GREEN,
    "Judges": "#7c3aed",
    "Deuteronomy": RED,
}


def figure_6() -> Path:
    rows = read_csv("chapter_stability_by_localization.csv")
    w, h = 1500, 1120
    img = Image.new("RGB", (w, h), PAPER)
    d = ImageDraw.Draw(img)
    title(d, "Figure 6. Chapter Stability Across Localization Cells", "Number of completed layer/α cells in which each chapter was rescued", w)
    left, right, top, bottom = 350, 1360, 180, 1020
    max_v = 43
    for tick in [0, 10, 20, 30, 40, 43]:
        x = int(left + tick / max_v * (right - left))
        d.line((x, top, x, bottom), fill="#e2e8f0", width=2)
        d.text((x, bottom + 18), str(tick), font=font(20), fill=GRAY, anchor="ma")
    row_h = (bottom - top) / len(rows)
    for i, r in enumerate(rows):
        y = int(top + row_h * i + row_h * 0.5)
        val = int(r["paired_rescue_count"])
        x = int(left + val / max_v * (right - left))
        color = BOOK_COLORS.get(r["book"], BLUE)
        d.text((left - 22, y), r["reference"], font=font(23, bold=i < 4), fill=NAVY, anchor="rm")
        d.rounded_rectangle((left, y - 13, x, y + 13), radius=9, fill=color)
        d.text((x + 14, y), str(val), font=font(22, bold=True), fill=NAVY, anchor="lm")
    out = FIG_DIR / "figure_6_chapter_stability.png"
    save_img(img, out)
    return out


def figure_7() -> Path:
    rows = read_csv("chapter_x_layer_alpha_rescue_matrix.csv")
    headers = [h for h in rows[0].keys() if h.startswith("L")]
    w, h = 2600, 1280
    img = Image.new("RGB", (w, h), PAPER)
    d = ImageDraw.Draw(img)
    title(d, "Figure 7. Chapter × Layer/α Rescue Matrix", "Blue cells mark a paired rescue for that chapter at that layer and steering strength", w)
    left, top = 390, 230
    cell_w, cell_h = 49, 47
    for j, head in enumerate(headers):
        x = left + j * cell_w + cell_w / 2
        layer, alpha = head[1:].split("_a")
        d.text((x, top - 67), f"L{layer}", font=font(13, bold=True), fill=NAVY, anchor="ma")
        d.text((x, top - 43), f"α{int(alpha)}", font=font(12), fill=GRAY, anchor="ma")
    last_layer = None
    for j, head in enumerate(headers):
        layer = head[1:].split("_a")[0]
        if last_layer is not None and layer != last_layer:
            x = left + j * cell_w
            d.line((x, top - 80, x, top + len(rows) * cell_h), fill="#94a3b8", width=3)
        last_layer = layer
    for i, r in enumerate(rows):
        y = top + i * cell_h
        d.text((left - 24, y + cell_h / 2), r["reference"], font=font(22, bold=i < 4), fill=NAVY, anchor="rm")
        for j, head in enumerate(headers):
            x = left + j * cell_w
            v = r[head] == "1"
            fill = BLUE if v else "#f8fafc"
            outline = "#cbd5e1"
            d.rectangle((x + 4, y + 4, x + cell_w - 4, y + cell_h - 4), fill=fill, outline=outline)
    out = FIG_DIR / "figure_7_rescue_matrix.png"
    save_img(img, out)
    return out


def build_figures() -> dict[str, Path]:
    for stale in FIG_DIR.glob("figure_*.png"):
        stale.unlink()
    figs = {
        "Figure 1": figure_1(),
        "Figure 2": figure_2(),
        "Figure 3": figure_3(),
        "Figure 4": figure_4(),
        "Figure 5": figure_5(),
        "Figure 6": figure_6(),
        "Figure 7": figure_7(),
    }
    build_contact_sheet(list(figs.values()))
    return figs


def build_contact_sheet(paths: list[Path]) -> None:
    thumb_w, thumb_h = 520, 360
    pad, label_h = 34, 42
    cols = 2
    rows = math.ceil(len(paths) / cols)
    img = Image.new("RGB", (cols * thumb_w + (cols + 1) * pad, rows * (thumb_h + label_h) + (rows + 1) * pad), PAPER)
    d = ImageDraw.Draw(img)
    for idx, path in enumerate(paths):
        thumb = Image.open(path).convert("RGB")
        thumb.thumbnail((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        col = idx % cols
        row = idx // cols
        x0 = pad + col * (thumb_w + pad)
        y0 = pad + row * (thumb_h + label_h + pad)
        x = x0 + (thumb_w - thumb.width) // 2
        img.paste(thumb, (x, y0))
        label = path.stem.replace("_", " ")
        d.text((x0, y0 + thumb_h + 8), label, font=font(24, bold=True), fill=NAVY)
    save_img(img, FIG_DIR / "figures_contact_sheet.png")


def table_1_rows() -> list[list[str]]:
    return [
        ["Stage", "Input", "Slice", "α", "Controls", "Output"],
        ["Book discovery", "66 biblical books", "Justice ratio, limit 10", "32", "control, +Scripture, -α, null", "19 book candidates"],
        ["Book confirmation", "19 book candidates", "Justice ratio, limit 40", "32", "same four-condition battery", "7 confirmed book sources"],
        ["Chapter discovery", "170 chapters from confirmed books", "Justice ratio, limit 10", "32", "same four-condition battery", "38 preliminary chapter hits"],
        ["Chapter confirmation", "38 chapter hits", "Justice ratio, limit 40", "32", "same four-condition battery", "16 confirmed chapter movers"],
        ["Layer/α localization", "16 confirmed chapter movers", "Justice ratio, limit 10", "16-96", "same four-condition battery", "43 completed behavior cells"],
    ]


def table_2_rows() -> list[list[str]]:
    conf = {r["target"]: r for r in read_csv("chapter_confirmation_l40_survivors.csv")}
    stab = read_csv("chapter_stability_by_localization.csv")
    motif_map = motif_rows()
    rows = [["Chapter", "Limit-40 confirmation", "Localization stability", "Best cell", "Biblical justice motif"]]
    for r in stab:
        c = conf[r["target"]]
        motif = motif_map.get(r["reference"], ("", ""))[0]
        rows.append([
            r["reference"],
            f"control {c['control_accuracy']}; +Scripture {c['positive_accuracy']}; Δ {c['positive_delta']}",
            f"{r['paired_rescue_count']}/43 cells",
            f"L{r['best_center_layer']} / α {r['best_alpha']}",
            motif,
        ])
    return rows


def motif_rows() -> dict[str, tuple[str, str]]:
    text = (WRITING_PACKET / "chapter_interpretive_motifs.md").read_text()
    out: dict[str, tuple[str, str]] = {}
    for line in text.splitlines():
        if not line.startswith("| ") or line.startswith("| ---") or "Confirmed chapter" in line:
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) >= 3:
            out[cells[0]] = (cells[1], cells[2])
    return out


def add_table(doc: Document, rows: list[list[str]], widths: list[float], font_size: float = 8.5) -> None:
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, txt in enumerate(row):
            cell = cells[ci]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, start=105, bottom=100, end=105)
            if ri == 0:
                shade_cell(cell, "EAF2FF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci in [2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(txt)
            run.font.name = "Times New Roman"
            run.font.size = Pt(font_size)
            if ri == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(23, 32, 51)
            else:
                run.font.color.rgb = RGBColor(31, 41, 55)
            cell.width = Inches(widths[ci])
    doc.add_paragraph()


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(71, 85, 105)


def add_inline_runs(paragraph, text: str) -> None:
    # Conservative inline markdown support for paper prose.
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            format_run(run)
        token = m.group(0)
        content = token.strip("*`")
        run = paragraph.add_run(content)
        format_run(run)
        if token.startswith("**"):
            run.bold = True
        elif token.startswith("*"):
            run.italic = True
        elif token.startswith("`"):
            run.font.name = "Courier New"
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        format_run(run)


def format_run(run) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(31, 41, 55)


def set_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True
    for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)]:
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(23, 32, 51)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "ScriptureVec Justice Paper"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].font.name = "Times New Roman"
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "Draft generated from completed ScriptureVec Justice result packet"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.runs[0].font.name = "Times New Roman"
    fp.runs[0].font.size = Pt(8)
    fp.runs[0].font.color.rgb = RGBColor(100, 116, 139)


def make_docx(md: str, figs: dict[str, Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    set_styles(doc)
    add_header_footer(doc)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            doc.add_paragraph()
            i += 1
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_runs(p, stripped[2:].strip())
        elif stripped.startswith("## "):
            doc.add_paragraph(clean_md_inline(stripped[3:].strip()), style="Heading 1")
        elif stripped.startswith("### "):
            doc.add_paragraph(clean_md_inline(stripped[4:].strip()), style="Heading 2")
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.right_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(clean_md_inline(stripped[2:].strip()))
            format_run(r)
            r.italic = True
            r.font.color.rgb = RGBColor(71, 85, 105)
        elif stripped.startswith("[**Figure"):
            match = re.search(r"Figure (\d+)", stripped)
            if match:
                fnum = f"Figure {match.group(1)}"
                img = figs[fnum]
                doc.add_picture(str(img), width=Inches(6.85))
                cap = clean_md_inline(stripped.strip("[]"))
                cap = cap.replace(" placeholder.", ".")
                add_caption(doc, cap)
        elif stripped.startswith("[**Table 1"):
            add_caption(doc, "Table 1. Study pipeline and decision rules.")
            add_table(doc, table_1_rows(), [1.15, 1.35, 1.2, 0.55, 1.45, 1.25], font_size=8.3)
        elif stripped.startswith("[**Table 2"):
            add_caption(doc, "Table 2. Sixteen confirmed chapter movers with confirmation and localization summaries.")
            add_table(doc, table_2_rows(), [0.95, 1.35, 0.9, 0.75, 2.95], font_size=7.6)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:].strip())
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_inline_runs(p, stripped)
        i += 1
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


def write_review() -> None:
    review = """# ScriptureVec Justice Paper Evaluation

## Overall Assessment

The draft summarizes the project well. It leads with the positive discovery, preserves the ambition of the claim, and correctly frames the study as a resolution pipeline: canon to book, book to chapter, chapter to layer and steering strength. It also does a good job resisting the weaker claim that Scripture in general boosts Justice; the paper is about specific scriptural loci whose vectors move the benchmark.

## Updates Applied

- Updated the localization grid to the completed expanded 43-cell packet.
- Reframed the localization result around three regimes: `L30/a96` for broadest rescue, `L28/a16` for efficient low-alpha uptake, and `L24/a32` for strongest mean movement.
- Added an alpha-trajectory figure showing that different layers respond differently to steering strength.
- Updated chapter-stability counts from the completed expanded grid.
- Updated the layer/alpha matrix to sixteen chapters by forty-three completed cells.
- Removed the discovery-funnel figure and renumbered the remaining seven figures.
- Moved the confirmed-chapter table into the chapter-stability section, where it connects confirmation evidence to localization behavior.
- Removed the old interpretive scaffold table because the revised Section 7 now carries that argument in prose.
- Rewrote Section 7.2 to spend less space summarizing commentary and more space explaining why the chapter vectors plausibly moved the Justice benchmark.
- Added a separate Hebrews 7 subsection because its Melchizedekian priesthood signal is one of the most informative chapter-level hits.
- Removed the stale missing-cell limitation and backfill-next-step language, replacing it with wider-slice confirmation of the regime-defining cells.

## Still Worth Checking Before Publication

- Commentary citations in Section 7 should be verified against the actual commentaries before final publication.
- The reference list still has placeholders for author first names, venues, and ICMI numbers.
- The theological sentence in the abstract about the canon being received with “sufficient fidelity” is powerful, but it should stay tied to the empirical intervention rather than becoming a broader claim about the model.
- Localization is still limit-10 screen evidence; the chapter confirmation findings are stronger because they are limit-40. The regime-defining cells should be retested on a wider slice before final publication.
"""
    REVIEW_OUT.write_text(review)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    refresh_localization_key_data()
    md = update_markdown(SOURCE_MD.read_text())
    MD_OUT.write_text(md)
    figs = build_figures()
    make_docx(md, figs)
    write_review()
    PAPER_MD_OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(MD_OUT, PAPER_MD_OUT)
    shutil.copy2(DOCX_OUT, PAPER_DOCX_OUT)
    shutil.copy2(REVIEW_OUT, PAPER_REVIEW_OUT)
    print(DOCX_OUT)
    print(MD_OUT)
    print(REVIEW_OUT)


if __name__ == "__main__":
    main()
